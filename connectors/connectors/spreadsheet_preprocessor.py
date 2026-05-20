from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from math import isfinite
from typing import Any


_SLUG_RE = re.compile(r"[^a-z0-9]+")


class SpreadsheetPreprocessorUnavailable(Exception):
    """Raised when a required spreadsheet parser dependency is missing."""


@dataclass(frozen=True, slots=True)
class PreprocessedTable:
    name: str
    label: str
    sheet_name: str
    columns: list[str]
    rows: list[list[Any]]
    source: str = "canonical_grid"
    metadata: dict[str, Any] = field(default_factory=dict)
    source_row_numbers: list[int] = field(default_factory=list)


def slug(value: str) -> str:
    text = _SLUG_RE.sub("_", value.lower()).strip("_")
    return text or "sheet"


def _json_safe_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool | int):
        return value
    if isinstance(value, float):
        return value if isfinite(value) else None
    text = str(value).strip()
    return text or None


def _row_width(values: list[list[Any]]) -> int:
    return max((len(row) for row in values), default=0)


def _has_any_value(values: list[list[Any]]) -> bool:
    return any(_json_safe_cell(cell) is not None for row in values for cell in row)


def _canonical_columns(width: int) -> list[str]:
    return ["row_number", *[f"column_{idx + 1}" for idx in range(width)]]


def _canonical_row(row_number: int, row: list[Any], width: int) -> list[Any]:
    cells = [
        _json_safe_cell(row[idx] if idx < len(row) else None)
        for idx in range(width)
    ]
    return [row_number, *cells]


def _is_blank_row(row: list[Any]) -> bool:
    return all(_json_safe_cell(cell) is None for cell in row)


def _non_empty_indexes(row: list[Any]) -> list[int]:
    return [idx for idx, cell in enumerate(row) if _json_safe_cell(cell) is not None]


def _looks_like_text_label(value: Any) -> bool:
    cell = _json_safe_cell(value)
    if cell is None or isinstance(cell, bool | int | float):
        return False
    if isinstance(cell, datetime | date | time):
        return False
    text = str(cell).strip()
    if not text or len(text) > 80:
        return False
    return any(ch.isalpha() for ch in text)


def _looks_numeric_or_temporal(value: Any) -> bool:
    cell = _json_safe_cell(value)
    if cell is None:
        return False
    if isinstance(cell, bool):
        return False
    if isinstance(cell, int | float | Decimal | datetime | date | time):
        return True
    text = str(cell).strip()
    if not text:
        return False
    if re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", text):
        return True
    if re.fullmatch(r"\d{1,4}[-/]\d{1,2}[-/]\d{1,4}", text):
        return True
    return False


def _looks_like_title_row(row: list[Any]) -> bool:
    indexes = _non_empty_indexes(row)
    return len(indexes) == 1 and _looks_like_text_label(row[indexes[0]])


def _next_non_blank_row(
    rows: list[tuple[int, list[Any]]],
    start_idx: int,
    *,
    max_distance: int = 3,
) -> tuple[int, list[Any]] | None:
    for idx in range(start_idx + 1, min(len(rows), start_idx + 1 + max_distance)):
        if not _is_blank_row(rows[idx][1]):
            return rows[idx]
    return None


def _header_score(rows: list[tuple[int, list[Any]]], idx: int) -> float:
    row = rows[idx][1]
    indexes = _non_empty_indexes(row)
    if len(indexes) < 2:
        return 0.0

    non_empty = len(indexes)
    text_count = sum(1 for col_idx in indexes if _looks_like_text_label(row[col_idx]))
    numeric_count = sum(1 for col_idx in indexes if _looks_numeric_or_temporal(row[col_idx]))
    if text_count == 0:
        return 0.0

    next_row = _next_non_blank_row(rows, idx)
    next_numeric = 0
    next_non_empty = 0
    if next_row is not None:
        next_indexes = _non_empty_indexes(next_row[1])
        next_non_empty = len(next_indexes)
        next_numeric = sum(
            1 for col_idx in next_indexes if _looks_numeric_or_temporal(next_row[1][col_idx])
        )

    score = text_count / non_empty
    score += min(non_empty, 8) / 8
    if next_non_empty >= max(1, min(2, non_empty)):
        score += 0.25
    if next_numeric > 0:
        score += 0.35
    if idx == 0 or _is_blank_row(rows[idx - 1][1]) or _looks_like_title_row(rows[idx - 1][1]):
        score += 0.2
    if numeric_count:
        score -= min(0.5, numeric_count / non_empty)
    return score


def _header_row_index(rows: list[tuple[int, list[Any]]]) -> tuple[int, float] | None:
    scored = [
        (idx, _header_score(rows, idx))
        for idx in range(len(rows))
    ]
    scored = [(idx, score) for idx, score in scored if score >= 1.25]
    if not scored:
        return None
    return max(scored, key=lambda item: item[1])


def _dedupe_columns(raw_columns: list[Any]) -> list[str]:
    columns: list[str] = []
    seen: dict[str, int] = {}
    for idx, value in enumerate(raw_columns):
        safe = _json_safe_cell(value)
        text = re.sub(r"\s+", " ", str(safe).strip()) if safe is not None else ""
        name = text[:80] if text else f"column_{idx + 1}"
        count = seen.get(name.lower(), 0) + 1
        seen[name.lower()] = count
        columns.append(name if count == 1 else f"{name} {count}")
    return columns


def _split_non_blank_regions(values: list[list[Any]]) -> list[list[tuple[int, list[Any]]]]:
    regions: list[list[tuple[int, list[Any]]]] = []
    current: list[tuple[int, list[Any]]] = []
    blank_run = 0
    for row_number, row in enumerate(values, start=1):
        if _is_blank_row(row):
            blank_run += 1
            if current and blank_run >= 2:
                regions.append(current)
                current = []
            continue
        blank_run = 0
        current.append((row_number, row))
    if current:
        regions.append(current)
    return regions


def _active_column_indexes(
    rows: list[tuple[int, list[Any]]],
    header_idx: int,
) -> list[int]:
    header_row = rows[header_idx][1]
    active = set(_non_empty_indexes(header_row))
    for _, row in rows[header_idx + 1:]:
        active.update(_non_empty_indexes(row))
    return sorted(active)


def _logical_table_from_region(
    sheet_name: str,
    rows: list[tuple[int, list[Any]]],
    *,
    source: str,
    metadata: dict[str, Any] | None,
) -> PreprocessedTable | None:
    header = _header_row_index(rows)
    if header is None:
        return None

    header_idx, score = header
    column_indexes = _active_column_indexes(rows, header_idx)
    if len(column_indexes) < 2:
        return None

    title_row_number: int | None = None
    title_label: str | None = None
    if header_idx > 0 and _looks_like_title_row(rows[header_idx - 1][1]):
        title_row_number, title_row = rows[header_idx - 1]
        title_index = _non_empty_indexes(title_row)[0]
        title_label = str(_json_safe_cell(title_row[title_index]) or "").strip()

    raw_columns = [
        rows[header_idx][1][col_idx] if col_idx < len(rows[header_idx][1]) else None
        for col_idx in column_indexes
    ]
    columns = _dedupe_columns(raw_columns)

    data_rows: list[list[Any]] = []
    source_row_numbers: list[int] = []
    header_normalized = [str(_json_safe_cell(cell) or "").strip().lower() for cell in raw_columns]
    for source_row_number, row in rows[header_idx + 1:]:
        selected = [
            _json_safe_cell(row[col_idx] if col_idx < len(row) else None)
            for col_idx in column_indexes
        ]
        if all(cell is None for cell in selected):
            continue
        selected_normalized = [str(cell or "").strip().lower() for cell in selected]
        if selected_normalized == header_normalized:
            continue
        data_rows.append(selected)
        source_row_numbers.append(source_row_number)

    if not data_rows:
        return None

    confidence = min(0.98, round(score / 2.4, 2))
    label = title_label or sheet_name
    table_metadata = {
        "format": "logical_table_v1",
        "sheet_name": sheet_name,
        "logical_label": label,
        "extraction_strategy": "deterministic_header_region",
        "confidence": confidence,
        "needs_agent_review": confidence < 0.72,
        "source_region": {
            "start_row": rows[0][0],
            "end_row": rows[-1][0],
            "header_row": rows[header_idx][0],
            "title_row": title_row_number,
            "columns": [idx + 1 for idx in column_indexes],
        },
        "raw_grid": {
            "grid_width": _row_width([row for _, row in rows]),
            "grid_height": len(rows),
        },
        **(metadata or {}),
    }
    return PreprocessedTable(
        name=slug(label),
        label=label,
        sheet_name=sheet_name,
        columns=columns,
        rows=data_rows,
        source=source,
        metadata=table_metadata,
        source_row_numbers=source_row_numbers,
    )


def _canonical_table(
    sheet_name: str,
    values: list[list[Any]],
    *,
    source: str,
    metadata: dict[str, Any] | None,
) -> PreprocessedTable | None:
    width = _row_width(values)
    if width == 0 or not _has_any_value(values):
        return None

    rows = [
        _canonical_row(row_number, row, width)
        for row_number, row in enumerate(values, start=1)
    ]
    table_metadata = {
        "format": "canonical_grid_v1",
        "sheet_name": sheet_name,
        "grid_width": width,
        "grid_height": len(values),
        "extraction_strategy": "raw_grid_fallback",
        "needs_agent_review": True,
        **(metadata or {}),
    }
    return PreprocessedTable(
        name=slug(sheet_name),
        label=sheet_name,
        sheet_name=sheet_name,
        columns=_canonical_columns(width),
        rows=rows,
        source=source,
        metadata=table_metadata,
        source_row_numbers=list(range(1, len(rows) + 1)),
    )


def _unique_table_names(tables: list[PreprocessedTable]) -> list[PreprocessedTable]:
    seen: dict[str, int] = {}
    unique: list[PreprocessedTable] = []
    for table in tables:
        base = table.name[:58] or "table"
        count = seen.get(base, 0) + 1
        seen[base] = count
        name = base if count == 1 else f"{base}_{count}"
        unique.append(
            PreprocessedTable(
                name=name,
                label=table.label,
                sheet_name=table.sheet_name,
                columns=table.columns,
                rows=table.rows,
                source=table.source,
                metadata=table.metadata,
                source_row_numbers=table.source_row_numbers,
            )
        )
    return unique


def preview_rows(
    table: PreprocessedTable,
    *,
    limit: int = 25,
    max_columns: int = 30,
) -> list[dict[str, Any]]:
    columns = table.columns[:max_columns]
    return [
        {
            column: row[idx] if idx < len(row) else None
            for idx, column in enumerate(columns)
        }
        for row in table.rows[:limit]
    ]


def preprocess_grid(
    sheet_name: str,
    values: list[list[Any]],
    *,
    source: str = "canonical_grid",
    metadata: dict[str, Any] | None = None,
) -> list[PreprocessedTable]:
    """Convert a worksheet into Baseflo-ready logical tables.

    Deterministic extraction handles the common spreadsheet mess:
    title rows, blank spacer rows, repeated headers, and human column names.
    When the shape is too ambiguous, Baseflo keeps a raw canonical grid and
    marks it for agent or human structure review.
    """
    if _row_width(values) == 0 or not _has_any_value(values):
        return []

    tables: list[PreprocessedTable] = []
    for region in _split_non_blank_regions(values):
        table = _logical_table_from_region(
            sheet_name,
            region,
            source=source,
            metadata=metadata,
        )
        if table is not None:
            tables.append(table)

    if tables:
        return _unique_table_names(tables)

    fallback = _canonical_table(sheet_name, values, source=source, metadata=metadata)
    return [fallback] if fallback is not None else []


def preprocess_xlsx_bytes(
    file_bytes: bytes,
    *,
    filename: str | None = None,
) -> list[PreprocessedTable]:
    """Convert every XLSX tab into Baseflo-ready logical tables."""
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - exercised when dependency is absent.
        raise SpreadsheetPreprocessorUnavailable(str(exc)) from exc

    workbook = load_workbook(
        io.BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )
    tables: list[PreprocessedTable] = []
    for worksheet in workbook.worksheets:
        values = [
            list(row)
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=worksheet.max_row,
                min_col=1,
                max_col=worksheet.max_column,
                values_only=True,
            )
        ]
        tables.extend(
            preprocess_grid(
                worksheet.title,
                values,
                source="openpyxl_canonical_grid",
                metadata={"filename": filename},
            )
        )
    return tables
